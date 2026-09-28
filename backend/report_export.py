"""Render analyst report data into downloadable PDF and DOCX files."""

from datetime import datetime
from html import escape
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak,
)


def _pretty(value):
    if value is None or value == "":
        return "—"
    if isinstance(value, (dict, list)):
        import json
        return json.dumps(value, indent=2, default=str)
    return str(value)


def _report_title(data):
    return "Alert Report" if data["report_type"] == "alerts" else "Incident Report"


def render_docx(data):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(45, 45, 55)

    document.add_heading("THREATHUNTER  /  SECURITY OPERATIONS", 2)
    document.add_heading(_report_title(data), 0)
    document.add_paragraph(f"Generated: {data['generated_at']}  |  Records: {data['record_count']}")
    document.add_heading("Summary", 1)
    document.add_paragraph(data["summary"])

    if data["report_type"] == "alerts":
        document.add_heading("Severity breakdown", 1)
        _docx_table(document, ["Severity", "Count"], [[key.title(), value] for key, value in data["severity_breakdown"].items()])
        document.add_heading("Alert types", 1)
        _docx_table(document, ["Detection type", "Count"], [[key, value] for key, value in data["alert_types"].items()])
        document.add_heading("Alerts", 1)
        for alert in data["alerts"]:
            document.add_heading(f"ALT-{alert['id']} · {alert['rule_name']}", 2)
            _docx_table(document, ["Timestamp", "Severity", "Status", "Source", "Destination"], [[
                _pretty(alert["timestamp"]), alert["severity"], alert["status"], alert["source_ip"], alert["destination_ip"],
            ]])
            document.add_paragraph(alert["description"] or "No description recorded.")
            document.add_paragraph("Recorded evidence", style="Heading 3")
            paragraph = document.add_paragraph(_pretty(alert["evidence"]))
            paragraph.style = "No Spacing"
    else:
        for incident in data["incidents"]:
            document.add_heading(f"INC-{incident['id']} · {incident['title']}", 1)
            document.add_paragraph(incident["summary"] or "No summary recorded.")
            _docx_table(document, ["Severity", "Status", "Created", "Updated"], [[
                incident["severity"], incident["status"], _pretty(incident["created_at"]), _pretty(incident["updated_at"]),
            ]])
            document.add_heading("Chronological alert timeline", 2)
            timeline = [["Time", "Rule / type", "Severity", "Source", "Destination"]]
            timeline.extend([[ _pretty(item["timestamp"]), item["rule_name"], item["severity"], item["source_ip"], item["destination_ip"]]
                             for item in incident["timeline"]])
            _docx_table(document, timeline[0], timeline[1:])
            document.add_heading("Related telemetry", 2)
            telemetry = [["Timestamp", "Event", "Source", "Destination", "Protocol", "Details"]]
            telemetry.extend([[
                _pretty(item["timestamp"]), item["event_type"], item["source_ip"], item["destination_ip"],
                item["protocol"], item["details"],
            ] for item in incident["related_telemetry"]])
            _docx_table(document, telemetry[0], telemetry[1:])
            document.add_heading("Indicators of compromise", 2)
            for label, key in (("IP addresses", "ip_addresses"), ("Domains", "domains"), ("URLs", "urls"), ("Destination ports", "destination_ports")):
                document.add_paragraph(f"{label}: {', '.join(map(str, incident['iocs'][key])) or 'None recorded'}")
            document.add_heading("MITRE ATT&CK techniques", 2)
            for technique in incident["mitre_techniques"]:
                document.add_paragraph(f"{technique['technique_id']} · {technique['name']} · {technique['tactic']}")
            if not incident["mitre_techniques"]:
                document.add_paragraph("No techniques associated.")
            document.add_heading("Investigation notes", 2)
            for note in incident["notes"]:
                document.add_paragraph(f"{_pretty(note['created_at'])}: {note['content']}")
            if not incident["notes"]:
                document.add_paragraph("No analyst notes recorded.")
            document.add_heading("Investigation actions", 2)
            for action in incident["actions"]:
                document.add_paragraph(f"{_pretty(action['created_at'])} · {action['action_type']}: {action['comment'] or 'No comment'}")
            if not incident["actions"]:
                document.add_paragraph("No investigation actions recorded.")
            document.add_heading("Associated alerts and evidence", 2)
            for alert in incident["alerts"]:
                document.add_heading(f"ALT-{alert['id']} · {alert['rule_name']}", 3)
                _docx_table(document, ["Time", "Severity", "Status", "Source", "Destination"], [[
                    _pretty(alert["timestamp"]), alert["severity"], alert["status"], alert["source_ip"], alert["destination_ip"],
                ]])
                document.add_paragraph(alert["description"] or "No description recorded.")
                document.add_paragraph("Evidence: " + _pretty(alert["evidence"]))

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _docx_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _pretty(value)
    document.add_paragraph()


def render_pdf(data):
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm, title=_report_title(data),
    )
    base = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#292633"), alignment=TA_LEFT, spaceAfter=7)
    heading = ParagraphStyle("ReportHeading", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#5e4bc6"), spaceBefore=12, spaceAfter=5)
    subheading = ParagraphStyle("ReportSubheading", parent=base["Heading3"], fontSize=10, leading=13, textColor=colors.HexColor("#363440"), spaceBefore=7, spaceAfter=3)
    body = ParagraphStyle("ReportBody", parent=base["BodyText"], fontSize=8, leading=11, textColor=colors.HexColor("#41404a"), spaceAfter=4, wordWrap="CJK")
    small = ParagraphStyle("ReportSmall", parent=body, fontSize=7, leading=9)
    evidence_style = ParagraphStyle("ReportEvidence", parent=small, fontName="Courier", backColor=colors.HexColor("#f5f4f8"), borderColor=colors.HexColor("#e4e1ed"), borderWidth=0.5, borderPadding=5, wordWrap="CJK")
    story = [Paragraph("THREATHUNTER  /  SECURITY OPERATIONS", small), Spacer(1, 3), Paragraph(_report_title(data), title),
             Paragraph(f"Generated: {escape(str(data['generated_at']))} &nbsp;&nbsp;|&nbsp;&nbsp; Records: {data['record_count']}", body),
             Paragraph("Summary", heading), Paragraph(escape(data["summary"]), body)]

    def add_table(headers, rows, widths=None):
        content = [[Paragraph(escape(str(item)), small) for item in headers]]
        content.extend([[Paragraph(escape(_pretty(item)).replace("\n", "<br/>"), small) for item in row] for row in rows])
        table = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0eef7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5143a8")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e2e0e8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        story.append(Spacer(1, 5))

    if data["report_type"] == "alerts":
        story.append(Paragraph("Severity breakdown", heading))
        add_table(["Severity", "Count"], [[key.title(), value] for key, value in data["severity_breakdown"].items()], [67 * mm, 35 * mm])
        story.append(Paragraph("Alert types", heading))
        add_table(["Detection type", "Count"], [[key, value] for key, value in data["alert_types"].items()], [105 * mm, 35 * mm])
        story.append(Paragraph("Alerts and recorded evidence", heading))
        for alert in data["alerts"]:
            story.extend([Paragraph(f"ALT-{alert['id']} · {escape(alert['rule_name'])}", subheading),
                          Paragraph(escape(alert["description"] or "No description recorded."), body)])
            add_table(["Timestamp", "Severity", "Status", "Source", "Destination"], [[
                _pretty(alert["timestamp"]), alert["severity"], alert["status"], alert["source_ip"], alert["destination_ip"],
            ]], [36 * mm, 25 * mm, 27 * mm, 41 * mm, 41 * mm])
            story.append(Paragraph("Recorded evidence", small))
            story.append(Paragraph(escape(_pretty(alert["evidence"])).replace("\n", "<br/>"), evidence_style))
            story.append(Spacer(1, 8))
    else:
        for incident in data["incidents"]:
            story.extend([Paragraph(f"INC-{incident['id']} · {escape(incident['title'])}", heading), Paragraph(escape(incident["summary"] or "No summary recorded."), body)])
            add_table(["Severity", "Status", "Created", "Updated"], [[incident["severity"], incident["status"], _pretty(incident["created_at"]), _pretty(incident["updated_at"])]], [28 * mm, 28 * mm, 67 * mm, 67 * mm])
            story.append(Paragraph("Chronological alert timeline", subheading))
            add_table(["Time", "Rule / type", "Severity", "Source", "Destination"], [[
                _pretty(item["timestamp"]), item["rule_name"], item["severity"], item["source_ip"], item["destination_ip"]
            ] for item in incident["timeline"]], [34 * mm, 43 * mm, 24 * mm, 43 * mm, 43 * mm])
            story.append(Paragraph("Related telemetry", subheading))
            add_table(["Timestamp", "Event", "Source", "Destination", "Protocol", "Details"], [[
                _pretty(item["timestamp"]), item["event_type"], item["source_ip"], item["destination_ip"],
                item["protocol"], item["details"],
            ] for item in incident["related_telemetry"]], [26 * mm, 18 * mm, 25 * mm, 25 * mm, 16 * mm, 55 * mm])
            story.append(Paragraph("Indicators of compromise", subheading))
            for label, key in (("IP addresses", "ip_addresses"), ("Domains", "domains"), ("URLs", "urls"), ("Destination ports", "destination_ports")):
                story.append(Paragraph(f"<b>{label}:</b> {escape(', '.join(map(str, incident['iocs'][key])) or 'None recorded')}", body))
            story.append(Paragraph("MITRE ATT&amp;CK techniques", subheading))
            techniques = incident["mitre_techniques"]
            story.append(Paragraph("<br/>".join(escape(f"{item['technique_id']} · {item['name']} · {item['tactic']}") for item in techniques) or "No techniques associated.", body))
            story.append(Paragraph("Investigation notes", subheading))
            story.append(Paragraph("<br/>".join(escape(f"{_pretty(item['created_at'])}: {item['content']}") for item in incident["notes"]) or "No analyst notes recorded.", body))
            story.append(Paragraph("Investigation actions", subheading))
            story.append(Paragraph("<br/>".join(escape(f"{_pretty(item['created_at'])} · {item['action_type']}: {item['comment'] or 'No comment'}") for item in incident["actions"]) or "No investigation actions recorded.", body))
            story.append(Paragraph("Associated alert evidence", subheading))
            for alert in incident["alerts"]:
                story.append(Paragraph(f"ALT-{alert['id']} · {escape(alert['rule_name'])}", small))
                story.append(Paragraph(escape(_pretty(alert["evidence"])).replace("\n", "<br/>"), evidence_style))
    document.build(story)
    output.seek(0)
    return output
